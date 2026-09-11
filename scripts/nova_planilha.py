"""Cria uma planilha nova (em branco, lista simples) para uma nova viagem e a
torna a planilha ativa para scripts/add_lancamento.py.

Uso:
    python3 scripts/nova_planilha.py "nome do arquivo sem extensão"

Exemplo:
    python3 scripts/nova_planilha.py "lancamentos em andamento"

Depois de rodar, use scripts/add_lancamento.py normalmente para ir adicionando
os lançamentos em ordem, e scripts/finalizar_planilha.py quando a viagem
terminar.
"""

import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

PLANILHA_DIR = Path(__file__).resolve().parent.parent / "planilha"
CURRENT_POINTER = PLANILHA_DIR / ".current"

HEADERS = [
    "DATES (JJ/MM/AAAA)",
    "DESIGNATION",
    "N° du justificatif",
    "MONTANT EN EUROS",
    "Classe",
]


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit('Uso: python3 nova_planilha.py "nome do arquivo sem extensão"')

    nome = sys.argv[1].strip()
    if not nome:
        sys.exit("Nome do arquivo não pode ser vazio.")

    filename = f"{nome}.xlsx"
    path = PLANILHA_DIR / filename

    if path.exists():
        sys.exit(f"Já existe um arquivo em {path}. Escolha outro nome.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dépenses"

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, title in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border

    widths = [20, 45, 20, 20, 15]
    for col, w in zip("ABCDE", widths):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    PLANILHA_DIR.mkdir(exist_ok=True)
    wb.save(path)
    CURRENT_POINTER.write_text(filename, encoding="utf-8")

    print(f"Planilha criada: {path.relative_to(PLANILHA_DIR.parent)}")
    print(f"Planilha ativa (usada por add_lancamento.py): {filename}")


if __name__ == "__main__":
    main()
