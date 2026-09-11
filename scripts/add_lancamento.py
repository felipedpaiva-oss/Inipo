"""Adiciona uma linha de lançamento à planilha planilha/Notes de frais Pouilles du Sud 06-09-26.xlsx.

Uso:
    python3 scripts/add_lancamento.py "DATA" "DESIGNATION" "N_JUSTIFICATIF" "MONTANT" "CLASSE"

Exemplo:
    python3 scripts/add_lancamento.py "06/09/2026" "Achat fournitures" "F2026-0002" 32.90 "Classe 2"

Os 5 argumentos devem seguir sempre esta ordem, que corresponde às colunas fixas da planilha:
DATES (JJ/MM/AAAA) | DESIGNATION | N° du justificatif | MONTANT EN EUROS | Classe
"""

import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Border, Side

XLSX_PATH = Path(__file__).resolve().parent.parent / "planilha" / "Notes de frais Pouilles du Sud 06-09-26.xlsx"


def main() -> None:
    if len(sys.argv) != 6:
        sys.exit(
            "Uso: python3 add_lancamento.py DATA DESIGNATION N_JUSTIFICATIF MONTANT CLASSE"
        )

    data, designation, justificatif, montant_raw, classe = sys.argv[1:6]

    try:
        montant = float(str(montant_raw).replace(",", "."))
    except ValueError:
        sys.exit(f"MONTANT inválido: {montant_raw!r} (use um número, ex: 45.50)")

    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb["Dépenses"]

    last_used_row = 2
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
        if row[0].value not in (None, ""):
            last_used_row = row[0].row
    next_row = last_used_row + 1
    values = [data, designation, justificatif, montant, classe]

    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    font = Font(name="Arial")

    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=next_row, column=col, value=value)
        cell.font = font
        cell.border = border
        if col == 4:
            cell.number_format = "#,##0.00"

    wb.save(XLSX_PATH)
    print(f"Linha {next_row} adicionada em {XLSX_PATH.relative_to(XLSX_PATH.parent.parent)}")


if __name__ == "__main__":
    main()
